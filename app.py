from flask import Flask, render_template, request, jsonify, redirect, url_for, session
from datetime import date, datetime, timedelta
import sqlite3, smtplib, json, os, calendar, io, csv
from openpyxl import load_workbook
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart

app = Flask(__name__)
app.secret_key = "reachai-tracker-secret-2026"

DB = os.path.join(os.path.dirname(__file__), "tracker.db")

SMTP_HOST = "smtp.gmail.com"
SMTP_PORT = 587
SMTP_USER = "parag@thinktoshareconnect.com"
SMTP_PASS = os.environ.get("SMTP_PASS", "")
ADMIN_EMAIL = "rahul.thinktoshare@gmail.com"

# Task definitions per employee
TASKS = {
    "ratul": [
        {"key": "whatsapp_sent",        "label": "WhatsApp Messages Sent",                    "type": "number",  "target": 35,  "unit": "messages", "time": "12:00–12:15"},
        {"key": "contacts_uploaded",    "label": "BNI Contacts Uploaded to ReachAI",           "type": "number",  "target": 130, "unit": "contacts", "time": "12:15–12:30"},
        {"key": "replies_checked",      "label": "Email + WhatsApp Replies Checked",            "type": "boolean", "target": 1,   "unit": "",         "time": "12:30–1:00"},
        {"key": "blog_url",             "label": "Getreached Blog Post Published",              "type": "url",     "target": 1,   "unit": "",         "time": "1:00–1:30"},
        {"key": "apollo_contacts",      "label": "Apollo.io Contacts Researched and Uploaded",  "type": "number",  "target": 150, "unit": "contacts", "time": "1:30–2:00"},
        {"key": "t2s_blog_url",         "label": "Think to Share Blog Post",                    "type": "url",     "target": 1,   "unit": "",         "time": "2:00–2:30"},
        {"key": "getreached_seo_notes", "label": "Getreached On Page SEO and Research",         "type": "text",    "target": 0,   "unit": "",         "time": "2:30–3:00"},
        {"key": "backlinks_built",      "label": "Backlinks for Getreached",                    "type": "number",  "target": 3,   "unit": "backlinks","time": "4:00–5:30"},
        {"key": "video_uploads",        "label": "Video Uploads",                               "type": "url",     "target": 1,   "unit": "",         "time": "5:30–6:30"},
    ],
    "debargha": [
        {"key": "linkedin_carousel",  "label": "LinkedIn Carousel Created",       "type": "url",    "target": 1,   "unit": "",         "time": "9:00–9:30"},
        {"key": "instagram_reel",     "label": "Instagram Reel Created",          "type": "url",    "target": 1,   "unit": "",         "time": "9:30–10:30"},
        {"key": "li_sales_nav",       "label": "LinkedIn Sales Navigator Researched","type": "number","target": 40, "unit": "contacts", "time": "10:30–11:00"},
        {"key": "li_connections",     "label": "LinkedIn Connection Requests Sent","type": "number", "target": 50,  "unit": "requests", "time": "11:00–12:00"},
        {"key": "li_dms",             "label": "LinkedIn DMs Sent",               "type": "number", "target": 60,  "unit": "DMs",      "time": "11:00–12:00"},
        {"key": "messages_checked",   "label": "LinkedIn + Instagram Replies Checked","type": "boolean","target": 1,"unit": "",        "time": "12:00–12:30"},
        {"key": "insta_dms",          "label": "Instagram DMs Sent",              "type": "number", "target": 50,  "unit": "DMs",      "time": "2:00–3:00"},
        {"key": "fb_reddit_posts",    "label": "Facebook Post",                   "type": "number", "target": 1,   "unit": "posts",    "time": "3:00–3:30"},
        {"key": "fb_reddit_urls",     "label": "Facebook Post URL",               "type": "url",    "target": 1,   "unit": "",         "time": "3:00–3:30"},
        {"key": "youtube_url",        "label": "YouTube Video",                   "type": "url",    "target": 1,   "unit": "",         "time": "3:30–4:30"},
        {"key": "eod_notes",          "label": "EOD Notes",                       "type": "text",   "target": 0,   "unit": "",         "time": "5:30–6:00"},
    ]
}

USERS = {
    "parag":    {"password": "parag123",    "role": "admin",    "name": "Parag Roy"},
    "ratul":    {"password": "ratul123",    "role": "employee", "name": "Ratul"},
    "debargha": {"password": "debargha123", "role": "employee", "name": "Debargha"},
}


def get_db():
    conn = sqlite3.connect(DB)
    conn.row_factory = sqlite3.Row
    return conn


UPLOAD_TASKS = {
    "ratul": {
        "whatsapp_sent":       {"label": "WhatsApp Contacts",            "accounts": ["WhatsApp"]},
        "contacts_uploaded":   {"label": "ReachAI Email Contacts",       "accounts": ["ReachAI Email"]},
        "apollo_contacts":     {"label": "Apollo Researched Contacts",   "accounts": ["Apollo.io"]},
        "backlinks_built":     {"label": "Backlinks",                    "accounts": ["Backlinks"]},
    },
    "debargha": {
        "li_connections": {"label": "LinkedIn Connection Contacts", "accounts": ["Parag LinkedIn", "Souranka LinkedIn"]},
        "li_dms":         {"label": "LinkedIn DM Contacts",         "accounts": ["Parag LinkedIn", "Souranka LinkedIn"]},
        "insta_dms":      {"label": "Instagram DM Contacts",        "accounts": ["Instagram Profile 1", "Instagram Profile 2"]},
    }
}

COL_MAP = {
    "name": ["name","full name","contact name","person","contact","first name","firstname","user"],
    "last_name": ["last name","lastname","surname"],
    "phone_number": ["phone","phone number","mobile","contact number","tel","mobile number","work phone","direct phone"],
    "email_id": ["email","email id","email address","e-mail","work email","corporate email"],
    "business_name": ["business name","company","organization","firm","business","company name","employer","account name"],
    "business_category": ["business category","industry","category","sector","type","vertical","niche"],
    "profile_url": ["profile url","linkedin url","instagram url","url","link","profile link","linkedin","person linkedin url","linkedin profile","backlink url","backlink","page url","website url"],
}

def _map_columns(headers):
    mapping = {}
    headers_lower = [h.strip().lower() if h else "" for h in headers]
    for field, aliases in COL_MAP.items():
        for alias in aliases:
            if alias in headers_lower:
                mapping[field] = headers_lower.index(alias)
                break
    # fallback: if no name field matched, use first non-empty column as name
    if "name" not in mapping and "last_name" not in mapping and headers_lower:
        for i, h in enumerate(headers_lower):
            if h:
                mapping["_fallback_name"] = i
                break
    return mapping

def _build_contact(get_fn, col_map):
    first = get_fn("name")
    last = get_fn("last_name")
    if first and last:
        name = first + " " + last
    elif first:
        name = first
    elif last:
        name = last
    else:
        name = get_fn("_fallback_name")
    return {
        "name": name,
        "phone_number": get_fn("phone_number"),
        "email_id": get_fn("email_id"),
        "business_name": get_fn("business_name"),
        "business_category": get_fn("business_category"),
        "profile_url": get_fn("profile_url"),
    }

def _parse_excel(file_bytes):
    wb = load_workbook(filename=io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(c).strip() if c else "" for c in rows[0]]
    col_map = _map_columns(headers)
    contacts = []
    for row in rows[1:]:
        if not any(row):
            continue
        def get_xl(field, r=row, m=col_map):
            idx = m.get(field)
            return str(r[idx]).strip() if idx is not None and idx < len(r) and r[idx] is not None else ""
        contacts.append(_build_contact(get_xl, col_map))
    return contacts

def _parse_csv(file_bytes):
    text = file_bytes.decode("utf-8", errors="ignore")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return []
    headers = [h.strip() for h in rows[0]]
    col_map = _map_columns(headers)
    contacts = []
    for row in rows[1:]:
        if not any(row):
            continue
        def get_csv(field, r=row, m=col_map):
            idx = m.get(field)
            return r[idx].strip() if idx is not None and idx < len(r) else ""
        contacts.append(_build_contact(get_csv, col_map))
    return contacts


def init_db():
    conn = get_db()
    conn.execute("""
        CREATE TABLE IF NOT EXISTS daily_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL,
            log_date TEXT NOT NULL,
            submitted INTEGER DEFAULT 0,
            submitted_at TEXT,
            UNIQUE(username, log_date)
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS task_entries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            log_id INTEGER NOT NULL,
            task_key TEXT NOT NULL,
            value_number INTEGER,
            value_url TEXT,
            value_bool INTEGER,
            value_text TEXT,
            updated_at TEXT,
            UNIQUE(log_id, task_key),
            FOREIGN KEY(log_id) REFERENCES daily_logs(id)
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS outreach_contacts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT,
            phone_number TEXT,
            email_id TEXT,
            business_name TEXT,
            business_category TEXT,
            profile_url TEXT,
            channel TEXT,
            account_used TEXT,
            contacted_by TEXT,
            contacted_date TEXT,
            log_id INTEGER,
            created_at TEXT
        )
    """)
    conn.commit()
    conn.close()


def get_or_create_log(username, log_date):
    conn = get_db()
    row = conn.execute(
        "SELECT id FROM daily_logs WHERE username=? AND log_date=?",
        (username, log_date)
    ).fetchone()
    if row:
        log_id = row["id"]
    else:
        cur = conn.execute(
            "INSERT INTO daily_logs (username, log_date) VALUES (?,?)",
            (username, log_date)
        )
        log_id = cur.lastrowid
        conn.commit()
    conn.close()
    return log_id


def get_log_data(username, log_date):
    conn = get_db()
    log = conn.execute(
        "SELECT * FROM daily_logs WHERE username=? AND log_date=?",
        (username, log_date)
    ).fetchone()
    if not log:
        conn.close()
        return None, {}
    entries = conn.execute(
        "SELECT * FROM task_entries WHERE log_id=?", (log["id"],)
    ).fetchall()
    conn.close()
    data = {}
    for e in entries:
        data[e["task_key"]] = dict(e)
    return dict(log), data


def send_missed_email(employee_name, missed_tasks, log_date):
    if not SMTP_PASS:
        return
    try:
        msg = MIMEMultipart("alternative")
        msg["Subject"] = f"⚠️ Missed Tasks — {employee_name} ({log_date})"
        msg["From"] = SMTP_USER
        msg["To"] = ADMIN_EMAIL

        rows = "".join(
            f"<tr><td style='padding:8px 12px;border-bottom:1px solid #eee;'>{t}</td></tr>"
            for t in missed_tasks
        )
        html = f"""
        <div style="font-family:sans-serif;max-width:520px;margin:0 auto;">
          <div style="background:#1e2a4a;color:#fff;padding:20px 24px;border-radius:8px 8px 0 0;">
            <h2 style="margin:0;font-size:18px;">ReachAI Activity Tracker</h2>
            <p style="margin:4px 0 0;opacity:.7;font-size:13px;">Daily Submission Report</p>
          </div>
          <div style="background:#fff8f0;border:1px solid #fbd9a0;padding:14px 24px;">
            <strong style="color:#b45309;">⚠️ {employee_name} submitted with {len(missed_tasks)} incomplete task(s) on {log_date}</strong>
          </div>
          <table style="width:100%;border-collapse:collapse;background:#fff;border:1px solid #e5e7eb;">
            <tr style="background:#f3f4f6;"><th style="padding:8px 12px;text-align:left;font-size:12px;color:#6b7280;">MISSED TASKS</th></tr>
            {rows}
          </table>
          <p style="font-size:12px;color:#9ca3af;padding:12px 0;">ReachAI Activity Tracker &nbsp;·&nbsp; Think To Share IT Solutions Pvt. Ltd.</p>
        </div>"""

        msg.attach(MIMEText(html, "html"))
        with smtplib.SMTP(SMTP_HOST, SMTP_PORT) as server:
            server.starttls()
            server.login(SMTP_USER, SMTP_PASS)
            server.sendmail(SMTP_USER, ADMIN_EMAIL, msg.as_string())
    except Exception as e:
        print(f"Email error: {e}")


@app.route("/")
def index():
    if "username" not in session:
        return redirect(url_for("login"))
    if session["role"] == "admin":
        return redirect(url_for("admin_dashboard"))
    return redirect(url_for("employee_dashboard"))


@app.route("/login", methods=["GET", "POST"])
def login():
    error = ""
    if request.method == "POST":
        username = request.form.get("username", "").strip().lower()
        password = request.form.get("password", "").strip()
        user = USERS.get(username)
        if user and user["password"] == password:
            session["username"] = username
            session["role"] = user["role"]
            session["name"] = user["name"]
            return redirect(url_for("index"))
        error = "Invalid username or password"
    return render_template("login.html", error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/dashboard")
def employee_dashboard():
    if "username" not in session or session["role"] != "employee":
        return redirect(url_for("login"))
    username = session["username"]
    today = date.today().isoformat()
    log_id = get_or_create_log(username, today)
    log, entries = get_log_data(username, today)
    tasks = TASKS.get(username, [])
    upload_tasks = UPLOAD_TASKS.get(username, {})
    return render_template("employee.html",
        username=username, name=session["name"],
        today=today, tasks=tasks, entries=entries,
        submitted=log.get("submitted", 0) if log else 0,
        upload_tasks=upload_tasks
    )


@app.route("/update-task", methods=["POST"])
def update_task():
    if "username" not in session or session["role"] != "employee":
        return jsonify({"error": "unauthorized"}), 401
    data = request.json
    username = session["username"]
    today = date.today().isoformat()
    log_id = get_or_create_log(username, today)

    task_key = data.get("key")
    task_type = data.get("type")
    now = datetime.now().isoformat()

    # backlinks saves number + urls text together
    if task_key == "backlinks_built" and task_type == "number":
        task_type = "backlinks"

    conn = get_db()
    existing = conn.execute(
        "SELECT id FROM task_entries WHERE log_id=? AND task_key=?",
        (log_id, task_key)
    ).fetchone()

    if task_type == "backlinks":
        val = int(data.get("value", 0))
        urls_text = data.get("urls_text", "")
        if existing:
            conn.execute("UPDATE task_entries SET value_number=?, value_text=?, updated_at=? WHERE log_id=? AND task_key=?",
                         (val, urls_text, now, log_id, task_key))
        else:
            conn.execute("INSERT INTO task_entries (log_id,task_key,value_number,value_text,updated_at) VALUES (?,?,?,?,?)",
                         (log_id, task_key, val, urls_text, now))
    elif task_type == "number":
        val = int(data.get("value", 0))
        if existing:
            conn.execute("UPDATE task_entries SET value_number=?, updated_at=? WHERE log_id=? AND task_key=?",
                         (val, now, log_id, task_key))
        else:
            conn.execute("INSERT INTO task_entries (log_id,task_key,value_number,updated_at) VALUES (?,?,?,?)",
                         (log_id, task_key, val, now))
    elif task_type == "url":
        val = data.get("value", "").strip()
        if existing:
            conn.execute("UPDATE task_entries SET value_url=?, updated_at=? WHERE log_id=? AND task_key=?",
                         (val, now, log_id, task_key))
        else:
            conn.execute("INSERT INTO task_entries (log_id,task_key,value_url,updated_at) VALUES (?,?,?,?)",
                         (log_id, task_key, val, now))
    elif task_type == "boolean":
        val = 1 if data.get("value") else 0
        if existing:
            conn.execute("UPDATE task_entries SET value_bool=?, updated_at=? WHERE log_id=? AND task_key=?",
                         (val, now, log_id, task_key))
        else:
            conn.execute("INSERT INTO task_entries (log_id,task_key,value_bool,updated_at) VALUES (?,?,?,?)",
                         (log_id, task_key, val, now))
    elif task_type == "text":
        val = data.get("value", "").strip()
        if existing:
            conn.execute("UPDATE task_entries SET value_text=?, updated_at=? WHERE log_id=? AND task_key=?",
                         (val, now, log_id, task_key))
        else:
            conn.execute("INSERT INTO task_entries (log_id,task_key,value_text,updated_at) VALUES (?,?,?,?)",
                         (log_id, task_key, val, now))

    conn.commit()
    conn.close()
    return jsonify({"status": "saved"})


@app.route("/submit-day", methods=["POST"])
def submit_day():
    if "username" not in session or session["role"] != "employee":
        return jsonify({"error": "unauthorized"}), 401
    username = session["username"]
    today = date.today().isoformat()
    log, entries = get_log_data(username, today)
    if not log:
        return jsonify({"error": "no log"}), 400

    tasks = TASKS.get(username, [])
    missed = []
    for t in tasks:
        if t["target"] == 0:
            continue
        e = entries.get(t["key"])
        if t["type"] == "number":
            val = e["value_number"] if e and e["value_number"] else 0
            if val < t["target"]:
                missed.append(f"{t['label']}: {val}/{t['target']} {t['unit']}")
        elif t["type"] == "url":
            val = e["value_url"] if e and e["value_url"] else ""
            if not val:
                missed.append(f"{t['label']}: No URL submitted")
        elif t["type"] == "boolean":
            val = e["value_bool"] if e and e["value_bool"] else 0
            if not val:
                missed.append(f"{t['label']}: Not completed")

    conn = get_db()
    conn.execute(
        "UPDATE daily_logs SET submitted=1, submitted_at=? WHERE id=?",
        (datetime.now().isoformat(), log["id"])
    )
    conn.commit()
    conn.close()

    if missed:
        send_missed_email(session["name"], missed, today)

    return jsonify({"status": "submitted", "missed": missed})


@app.route("/admin")
def admin_dashboard():
    if "username" not in session or session["role"] != "admin":
        return redirect(url_for("login"))
    emp = request.args.get("emp", "ratul")
    period = request.args.get("period", "today")
    today = date.today()

    if period == "week":
        start = today - timedelta(days=6)
    elif period == "month":
        start = today.replace(day=1)
    else:
        start = today

    conn = get_db()
    rows = conn.execute("""
        SELECT dl.id, dl.log_date, dl.submitted,
               te.task_key, te.value_number, te.value_bool, te.value_url, te.value_text
        FROM daily_logs dl
        LEFT JOIN task_entries te ON te.log_id = dl.id
        WHERE dl.username=? AND dl.log_date BETWEEN ? AND ?
        ORDER BY dl.log_date DESC
    """, (emp, start.isoformat(), today.isoformat())).fetchall()

    task_totals = {}
    task_urls = {}
    daily = {}
    days_submitted = 0
    days_seen = set()

    for r in rows:
        d = r["log_date"]
        if d not in daily:
            daily[d] = {"submitted": r["submitted"], "tasks": {}}
            days_seen.add(d)
            if r["submitted"]:
                days_submitted += 1
        if r["task_key"]:
            daily[d]["tasks"][r["task_key"]] = {
                "number": r["value_number"], "bool": r["value_bool"],
                "url": r["value_url"], "text": r["value_text"]
            }
            if r["value_number"] is not None:
                task_totals[r["task_key"]] = task_totals.get(r["task_key"], 0) + r["value_number"]
            if r["value_url"]:
                task_urls.setdefault(r["task_key"], []).append({"date": d, "url": r["value_url"]})
            if r["task_key"] == "backlinks_built" and r["value_text"]:
                for u in r["value_text"].split('\n'):
                    u = u.strip()
                    if u:
                        task_urls.setdefault("backlinks_built", []).append({"date": d, "url": u})

    contacts_rows = conn.execute("""
        SELECT channel, account_used, COUNT(*) as cnt
        FROM outreach_contacts
        WHERE contacted_by=? AND contacted_date BETWEEN ? AND ?
        GROUP BY channel, account_used ORDER BY cnt DESC
    """, (emp, start.isoformat(), today.isoformat())).fetchall()

    upload_actual_rows = conn.execute("""
        SELECT channel, COUNT(*) as cnt
        FROM outreach_contacts
        WHERE contacted_by=? AND contacted_date BETWEEN ? AND ?
        GROUP BY channel
    """, (emp, start.isoformat(), today.isoformat())).fetchall()
    conn.close()

    contacts_by_channel = {}
    for c in contacts_rows:
        ch = c["channel"]
        contacts_by_channel.setdefault(ch, {"total": 0, "accounts": []})
        contacts_by_channel[ch]["total"] += c["cnt"]
        contacts_by_channel[ch]["accounts"].append({"name": c["account_used"], "cnt": c["cnt"]})

    upload_actuals = {r["channel"]: r["cnt"] for r in upload_actual_rows}

    period_labels = {"today": "Today", "week": "Last 7 Days", "month": "This Month"}
    return render_template("admin.html",
        selected_emp=emp, selected_period=period,
        period_label=period_labels[period],
        today=today.isoformat(), start=start.isoformat(),
        task_totals=task_totals, task_urls=task_urls,
        contacts_by_channel=contacts_by_channel,
        upload_actuals=upload_actuals,
        daily=daily, days_submitted=days_submitted,
        days_logged=len(days_seen),
        tasks=TASKS.get(emp, []),
        emp_name=USERS[emp]["name"],
        users=USERS,
    )


@app.route("/admin/history")
def admin_history():
    if "username" not in session or session["role"] != "admin":
        return redirect(url_for("login"))
    employee = request.args.get("emp", "ratul")
    conn = get_db()
    logs = conn.execute(
        "SELECT * FROM daily_logs WHERE username=? ORDER BY log_date DESC LIMIT 30",
        (employee,)
    ).fetchall()
    history = []
    for log in logs:
        entries_raw = conn.execute(
            "SELECT * FROM task_entries WHERE log_id=?", (log["id"],)
        ).fetchall()
        entries = {e["task_key"]: dict(e) for e in entries_raw}
        tasks = TASKS.get(employee, [])
        completed = sum(1 for t in tasks if t["target"] > 0 and _is_done(t, entries))
        total = sum(1 for t in tasks if t["target"] > 0)
        history.append({
            "date": log["log_date"],
            "submitted": log["submitted"],
            "completed": completed,
            "total": total,
            "pct": round(completed / total * 100) if total else 0
        })
    conn.close()
    return render_template("history.html", history=history, employee=employee,
                           emp_name=USERS[employee]["name"])


def _is_done(task, entries):
    e = entries.get(task["key"])
    if task["type"] == "number":
        return (e["value_number"] if e and e["value_number"] else 0) >= task["target"]
    elif task["type"] == "url":
        return bool(e["value_url"] if e and e["value_url"] else "")
    elif task["type"] == "boolean":
        return bool(e["value_bool"] if e and e["value_bool"] else 0)
    return True


@app.route("/upload-contacts", methods=["POST"])
def upload_contacts():
    if "username" not in session or session["role"] != "employee":
        return jsonify({"error": "unauthorized"}), 401
    username = session["username"]
    today = date.today().isoformat()
    log_id = get_or_create_log(username, today)

    task_key = request.form.get("task_key")
    account_used = request.form.get("account_used", "")
    channel = request.form.get("channel", task_key)
    file = request.files.get("file")

    if not file:
        return jsonify({"error": "No file uploaded"}), 400

    file_bytes = file.read()
    filename = file.filename.lower()

    try:
        if filename.endswith(".xlsx") or filename.endswith(".xls"):
            contacts = _parse_excel(file_bytes)
        elif filename.endswith(".csv"):
            contacts = _parse_csv(file_bytes)
        else:
            return jsonify({"error": "Upload .xlsx or .csv file"}), 400
    except Exception as e:
        return jsonify({"error": f"Parse error: {str(e)}"}), 400
    if not contacts:
        return jsonify({"inserted": 0, "skipped": 0, "total": 0, "status": "ok"})

    now = datetime.now().isoformat()
    conn = get_db()
    inserted = 0
    skipped = 0
    for c in contacts:
        if not c.get("name") and not c.get("profile_url") and not c.get("email_id") and not c.get("phone_number"):
            continue
        # check duplicate — same person same channel same employee
        existing = conn.execute(
            """SELECT id FROM outreach_contacts
               WHERE contacted_by=? AND channel=? AND (
                 (profile_url!='' AND profile_url=?) OR
                 (email_id!='' AND email_id=?) OR
                 (phone_number!='' AND phone_number=?)
               )""",
            (username, channel, c["profile_url"], c["email_id"], c["phone_number"])
        ).fetchone()
        if existing:
            skipped += 1
            continue
        conn.execute(
            """INSERT INTO outreach_contacts
               (name,phone_number,email_id,business_name,business_category,
                profile_url,channel,account_used,contacted_by,contacted_date,log_id,created_at)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
            (c["name"], c["phone_number"], c["email_id"], c["business_name"],
             c["business_category"], c["profile_url"], channel, account_used,
             username, today, log_id, now)
        )
        inserted += 1
    conn.commit()
    conn.close()
    return jsonify({"status": "ok", "inserted": inserted, "skipped": skipped, "total": len(contacts)})


@app.route("/contacts")
def contacts_history():
    if "username" not in session:
        return redirect(url_for("login"))
    is_admin = session["role"] == "admin"
    if session["role"] == "employee":
        username = session["username"]
    else:
        username = request.args.get("emp", "all")
    channel = request.args.get("channel", "")
    search = request.args.get("q", "").strip().lower()
    conn = get_db()
    if is_admin and username == "all":
        query = "SELECT * FROM outreach_contacts WHERE 1=1"
        params = []
    else:
        query = "SELECT * FROM outreach_contacts WHERE contacted_by=?"
        params = [username]
    if channel:
        query += " AND channel=?"
        params.append(channel)
    query += " ORDER BY contacted_date DESC, id DESC LIMIT 1000"
    rows = conn.execute(query, params).fetchall()
    conn.close()

    if search:
        rows = [r for r in rows if search in (r["name"] or "").lower()
                or search in (r["business_name"] or "").lower()
                or search in (r["email_id"] or "").lower()
                or search in (r["profile_url"] or "").lower()]

    channels = list(set(r["channel"] for r in rows)) if rows else []
    emp_name = "All Employees" if (is_admin and username == "all") else USERS.get(username, {}).get("name", username)
    upload_tasks = UPLOAD_TASKS.get(username, {})
    return render_template("contacts.html",
        rows=rows, channels=channels, selected_channel=channel,
        search=search, emp_name=emp_name, username=username,
        is_admin=is_admin, upload_tasks=upload_tasks)


@app.route("/reopen-day", methods=["POST"])
def reopen_day():
    if "username" not in session or session["role"] != "employee":
        return jsonify({"error": "unauthorized"}), 401
    username = session["username"]
    today = date.today().isoformat()
    conn = get_db()
    conn.execute(
        "UPDATE daily_logs SET submitted=0, submitted_at=NULL WHERE username=? AND log_date=?",
        (username, today)
    )
    conn.commit()
    conn.close()
    return jsonify({"status": "reopened"})


def _get_range_data(username, start_date, end_date):
    conn = get_db()
    logs = conn.execute(
        "SELECT * FROM daily_logs WHERE username=? AND log_date>=? AND log_date<=? ORDER BY log_date",
        (username, start_date, end_date)
    ).fetchall()
    result = []
    for log in logs:
        entries_raw = conn.execute(
            "SELECT * FROM task_entries WHERE log_id=?", (log["id"],)
        ).fetchall()
        entries = {e["task_key"]: dict(e) for e in entries_raw}
        tasks = TASKS.get(username, [])
        completed = sum(1 for t in tasks if t["target"] > 0 and _is_done(t, entries))
        total = sum(1 for t in tasks if t["target"] > 0)
        # numeric task totals
        num_totals = {}
        for t in tasks:
            if t["type"] == "number":
                e = entries.get(t["key"], {})
                num_totals[t["key"]] = e.get("value_number") or 0
        result.append({
            "date": log["log_date"],
            "submitted": log["submitted"],
            "completed": completed,
            "total": total,
            "pct": round(completed / total * 100) if total else 0,
            "num_totals": num_totals,
            "entries": entries
        })
    conn.close()
    return result


@app.route("/admin/urls")
def admin_urls():
    if "username" not in session or session["role"] != "admin":
        return redirect(url_for("login"))
    conn = get_db()
    # URL-type tasks to show
    URL_TASKS = {
        "backlinks_built": "Backlinks",
        "blog_url": "Blog Post",
        "fb_reddit_urls": "Facebook/Reddit",
        "youtube_url": "YouTube Video",
        "li_connections": "LinkedIn Carousel",
        "instagram_reel": "Instagram Reel",
        "li_carousel": "LinkedIn Carousel",
    }
    rows = conn.execute("""
        SELECT dl.username, dl.log_date, te.task_key,
               te.value_url, te.value_text
        FROM task_entries te
        JOIN daily_logs dl ON te.log_id = dl.id
        WHERE te.task_key IN ({})
          AND (te.value_url IS NOT NULL AND te.value_url != ''
               OR te.value_text IS NOT NULL AND te.value_text != '')
        ORDER BY dl.log_date DESC, dl.username
    """.format(','.join('?' * len(URL_TASKS))), list(URL_TASKS.keys())).fetchall()
    conn.close()
    entries = []
    for r in rows:
        if r["task_key"] == "backlinks_built" and r["value_text"]:
            for url in r["value_text"].split('\n'):
                url = url.strip()
                if url:
                    entries.append({"date": r["log_date"], "username": r["username"],
                                    "label": "Backlink", "url": url})
        elif r["value_url"]:
            entries.append({"date": r["log_date"], "username": r["username"],
                            "label": URL_TASKS.get(r["task_key"], r["task_key"]),
                            "url": r["value_url"]})
    return render_template("admin_urls.html", entries=entries,
                           emp_names={u: USERS[u]["name"] for u in USERS})


@app.route("/admin/weekly")
def admin_weekly():
    if "username" not in session or session["role"] != "admin":
        return redirect(url_for("login"))
    today = date.today()
    start = today - timedelta(days=6)
    employees = ["ratul", "debargha"]
    data = {}
    for emp in employees:
        rows = _get_range_data(emp, start.isoformat(), today.isoformat())
        tasks = TASKS.get(emp, [])
        num_tasks = [t for t in tasks if t["type"] == "number" and t["target"] > 0]
        totals = {t["key"]: sum(r["num_totals"].get(t["key"], 0) for r in rows) for t in num_tasks}
        avg_pct = round(sum(r["pct"] for r in rows) / len(rows)) if rows else 0
        data[emp] = {
            "name": USERS[emp]["name"],
            "rows": rows,
            "num_tasks": num_tasks,
            "totals": totals,
            "avg_pct": avg_pct,
            "days_submitted": sum(1 for r in rows if r["submitted"])
        }
    date_range = [(start + timedelta(days=i)).isoformat() for i in range(7)]
    return render_template("weekly.html", data=data, date_range=date_range,
                           start=start.isoformat(), end=today.isoformat())


@app.route("/admin/monthly")
def admin_monthly():
    if "username" not in session or session["role"] != "admin":
        return redirect(url_for("login"))
    today = date.today()
    start = date(today.year, today.month, 1)
    employees = ["ratul", "debargha"]
    data = {}
    for emp in employees:
        rows = _get_range_data(emp, start.isoformat(), today.isoformat())
        tasks = TASKS.get(emp, [])
        num_tasks = [t for t in tasks if t["type"] == "number" and t["target"] > 0]
        totals = {t["key"]: sum(r["num_totals"].get(t["key"], 0) for r in rows) for t in num_tasks}
        targets_monthly = {t["key"]: t["target"] * len(rows) for t in num_tasks}
        avg_pct = round(sum(r["pct"] for r in rows) / len(rows)) if rows else 0
        # week buckets
        weeks = {}
        for r in rows:
            d = date.fromisoformat(r["date"])
            wk = f"Week {((d.day - 1) // 7) + 1}"
            if wk not in weeks:
                weeks[wk] = {"days": 0, "pct_sum": 0}
            weeks[wk]["days"] += 1
            weeks[wk]["pct_sum"] += r["pct"]
        week_avgs = {wk: round(v["pct_sum"] / v["days"]) for wk, v in weeks.items()}
        data[emp] = {
            "name": USERS[emp]["name"],
            "rows": rows,
            "num_tasks": num_tasks,
            "totals": totals,
            "targets_monthly": targets_monthly,
            "avg_pct": avg_pct,
            "days_submitted": sum(1 for r in rows if r["submitted"]),
            "total_days": len(rows),
            "week_avgs": week_avgs
        }
    return render_template("monthly.html", data=data,
                           month=today.strftime("%B %Y"), today=today.isoformat())


init_db()

if __name__ == "__main__":
    app.run(debug=True, port=8081)
